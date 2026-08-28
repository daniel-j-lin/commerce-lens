import csv
import sqlite3
from datetime import date
from decimal import Decimal, ROUND_DOWN, ROUND_HALF_EVEN, ROUND_UP, getcontext, localcontext
from pathlib import Path

import duckdb
import pytest
from openpyxl import Workbook

from commerce_lens.canonical import (
    CanonicalizationRequest,
    EligibilityMode,
    EligibilityState,
    EligibilityValueMapping,
    canonicalize_dataset,
    identity_mapping,
)
from commerce_lens.contracts.common import (
    FailureDetail,
    FailureStage,
    GroupingDimension,
    MetricState,
    PeriodDefinition,
    ScopeDefinition,
    ScopeFilter,
    SourceType,
)
from commerce_lens.contracts.evidence import MetricReference
from commerce_lens.contracts.plans import PlanMetricNode
from commerce_lens.contracts.requests import AnalysisRequest
from commerce_lens.contracts.sufficiency import DataSufficiencyResult, MetricEligibility, SufficiencyState
from commerce_lens.engine import MetricExecutionError, execute_plan
from commerce_lens.engine.execution import _revenue_change_dependency_results
from commerce_lens.engine.plan_builder import build_execution_plan
from commerce_lens.engine.populations import population_fingerprint, population_id_for_fingerprint
from commerce_lens.intake.registry import DatasetRegistry
from commerce_lens.contracts.execution import ExecutedResult
from commerce_lens.evidence.identifiers import sha256_file
from commerce_lens.metrics import (
    AOV_EXECUTION_IMPLEMENTATION_REF,
    EXECUTION_NOT_IMPLEMENTED_REF,
    METRIC_DEFINITION_VERSION,
    METRIC_REGISTRY_VERSION,
    ORDERS_EXECUTION_IMPLEMENTATION_REF,
    REVENUE_CHANGE_EXECUTION_IMPLEMENTATION_REF,
    REVENUE_EXECUTION_IMPLEMENTATION_REF,
)
from commerce_lens.persistence.artifact_store import ArtifactStore
from commerce_lens.persistence.metadata_store import MetadataStore


def test_revenue_executes_single_and_multiple_eligible_lines_with_decimal_authority(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(
        tmp_path,
        ("revenue",),
        [
            _row(order_id="o1", order_line_id="l1", order_date="2026-01-01", line_revenue="10.10"),
            _row(order_id="o2", order_line_id="l1", order_date="2026-01-02", line_revenue="0"),
            _row(order_id="o3", order_line_id="l1", order_date="2026-01-03", line_revenue="99.99"),
            _row(
                order_id="o4",
                order_line_id="l1",
                order_date="2026-01-01",
                line_revenue="100.00",
                eligibility_status="cancelled",
            ),
        ],
    )

    outcome = execute_plan(plan, canonical, store)

    baseline = _result(outcome, "revenue", "baseline")
    comparison = _result(outcome, "revenue", "comparison")
    assert baseline.value == Decimal("10.10")
    assert comparison.value == Decimal("99.99")
    assert isinstance(baseline.value, Decimal)
    assert not isinstance(baseline.value, float)


def test_revenue_uses_line_revenue_not_quantity_times_unit_price(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(
        tmp_path,
        ("revenue",),
        [_row(quantity="10", line_revenue="7.25", unit_price="99.99")],
    )

    result = _result(execute_plan(plan, canonical, store), "revenue", "baseline")

    assert result.value == Decimal("7.25")


def test_revenue_preserves_high_scale_decimal_and_repeatability(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(
        tmp_path,
        ("revenue",),
        [
            _row(order_id="o1", order_line_id="l1", line_revenue="10.123456789123"),
            _row(order_id="o1", order_line_id="l2", line_revenue="0.000000000007"),
        ],
    )

    first = execute_plan(plan, canonical, store)
    second = execute_plan(plan, canonical, store)

    assert _result(first, "revenue", "baseline").value == Decimal("10.123456789130")
    assert _result(first, "revenue", "baseline").value == _result(second, "revenue", "baseline").value
    assert _result(first, "revenue", "baseline").result_id != _result(second, "revenue", "baseline").result_id
    assert _result(first, "revenue", "baseline").result_fingerprint == _result(
        second,
        "revenue",
        "baseline",
    ).result_fingerprint


def test_repeated_equivalent_execution_uses_unique_event_ids_and_ordered_timestamps(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(tmp_path, ("revenue",), [_row()])

    first = execute_plan(plan, canonical, store)
    second = execute_plan(plan, canonical, store)
    first_record = _record(first, "revenue", "baseline")
    second_record = _record(second, "revenue", "baseline")
    first_result = _result(first, "revenue", "baseline")
    second_result = _result(second, "revenue", "baseline")

    assert first_record.execution_id != second_record.execution_id
    assert first_result.result_id != second_result.result_id
    assert first_result.value == second_result.value == Decimal("10.00")
    assert first_result.result_fingerprint == second_result.result_fingerprint
    assert first_result.execution_id == first_record.execution_id
    assert second_result.execution_id == second_record.execution_id
    assert first_record.started_at <= first_record.ended_at
    assert second_record.started_at <= second_record.ended_at


def test_orders_count_distinct_eligible_orders_not_lines(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(
        tmp_path,
        ("orders",),
        [
            _row(order_id="o1", order_line_id="l1", line_revenue="10.00"),
            _row(order_id="o1", order_line_id="l2", line_revenue="5.00"),
            _row(order_id="o2", order_line_id="l1", line_revenue="0"),
            _row(order_id="o3", order_line_id="l1", eligibility_status="cancelled"),
            _row(order_id="o4", order_line_id="l1", order_date="2026-01-03"),
        ],
    )

    outcome = execute_plan(plan, canonical, store)

    baseline = _result(outcome, "orders", "baseline")
    comparison = _result(outcome, "orders", "comparison")
    assert baseline.value == 2
    assert comparison.value == 1
    assert isinstance(baseline.value, int)
    assert not isinstance(baseline.value, float)


def test_aov_uses_executed_revenue_and_orders_dependencies(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(
        tmp_path,
        ("aov",),
        [
            _row(order_id="o1", order_line_id="l1", line_revenue="10.00"),
            _row(order_id="o1", order_line_id="l2", line_revenue="5.00"),
            _row(order_id="o2", order_line_id="l1", line_revenue="15.00"),
            _row(order_id="o3", order_line_id="l1", order_date="2026-01-03", line_revenue="10.00"),
        ],
    )

    outcome = execute_plan(plan, canonical, store)
    baseline = _result(outcome, "aov", "baseline")
    comparison = _result(outcome, "aov", "comparison")
    baseline_record = _record(outcome, "aov", "baseline")

    assert baseline.value == Decimal("15.00")
    assert comparison.value == Decimal("10.00")
    assert isinstance(baseline.value, Decimal)
    assert _result(outcome, "revenue", "baseline").currency == "USD"
    assert baseline.currency == "USD"
    assert baseline_record.operation["method"] == "python_decimal_dependency_arithmetic"
    assert baseline_record.operation["calculation_policy"] == {
        "calculation_policy_id": "p4_aov_decimal_calculation_policy_v1",
        "precision": 38,
        "rounding": "ROUND_HALF_EVEN",
    }
    assert "revenue_result_ref" in baseline_record.operation
    assert "orders_result_ref" in baseline_record.operation


def test_aov_zero_revenue_with_orders_is_decimal_zero(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(
        tmp_path,
        ("aov",),
        [_row(order_id="o1", order_line_id="l1", line_revenue="0")],
    )

    result = _result(execute_plan(plan, canonical, store), "aov", "baseline")

    assert result.value == Decimal("0")
    assert isinstance(result.value, Decimal)


def test_aov_orders_zero_is_undefined_not_execution_failure(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(
        tmp_path,
        ("aov",),
        [_row(order_id="o1", order_line_id="l1", eligibility_status="cancelled")],
    )

    outcome = execute_plan(plan, canonical, store)
    result = _result(outcome, "aov", "baseline")
    record = _record(outcome, "aov", "baseline")

    assert result.value is None
    assert result.metric_state is MetricState.UNDEFINED
    assert result.undefined_reason == "orders_equals_zero"
    assert result.execution_status.value == "completed"
    assert record.status.value == "completed"


def test_explicit_currency_empty_physical_baseline_executes_zero_and_persists(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(
        tmp_path,
        ("aov",),
        [_row(order_id="o2", order_date="2026-01-03", line_revenue="12.00", currency="USD")],
        filename="explicit_currency_empty_baseline.csv",
        scope=ScopeDefinition(
            scope_id="usd_only",
            filters=(ScopeFilter(field="currency", operator="equals", value="USD"),),
        ),
    )
    metadata_store = MetadataStore(tmp_path / "explicit_empty_registry.sqlite")

    outcome = execute_plan(plan, canonical, store, metadata_store)
    revenue = _result(outcome, "revenue", "baseline")
    orders = _result(outcome, "orders", "baseline")
    aov = _result(outcome, "aov", "baseline")
    revenue_record = _record(outcome, "revenue", "baseline")
    orders_record = _record(outcome, "orders", "baseline")
    aov_record = _record(outcome, "aov", "baseline")

    assert revenue.value == Decimal("0")
    assert revenue.metric_state is MetricState.VALID
    assert revenue.currency == "USD"
    assert revenue_record.resolved_currency == "USD"
    assert revenue_record.eligible_input_row_count == 0
    assert orders.value == 0
    assert orders.metric_state is MetricState.VALID
    assert orders.currency is None
    assert orders_record.eligible_input_row_count == 0
    assert aov.value is None
    assert aov.metric_state is MetricState.UNDEFINED
    assert aov.undefined_reason == "orders_equals_zero"
    assert aov.currency == "USD"
    assert aov_record.resolved_currency == "USD"
    assert aov_record.eligible_input_row_count == 0

    restored_revenue = ExecutedResult.model_validate_json(
        store.safe_path(revenue_record.output_artifacts[0].path).read_text(encoding="utf-8")
    )
    restored_aov = ExecutedResult.model_validate_json(
        store.safe_path(aov_record.output_artifacts[0].path).read_text(encoding="utf-8")
    )
    assert restored_revenue == revenue
    assert restored_revenue.value == Decimal("0")
    assert isinstance(restored_revenue.value, Decimal)
    assert restored_revenue.metric_state is MetricState.VALID
    assert restored_revenue.currency == "USD"
    assert restored_aov == aov
    assert restored_aov.value is None
    assert restored_aov.metric_state is MetricState.UNDEFINED
    assert restored_aov.undefined_reason == "orders_equals_zero"
    assert restored_aov.currency == "USD"


def test_phase2_single_currency_empty_physical_baseline_uses_dataset_authority(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(
        tmp_path,
        ("aov",),
        [_row(order_id="o2", order_date="2026-01-03", line_revenue="12.00", currency="USD")],
        filename="phase2_currency_empty_baseline.csv",
    )

    outcome = execute_plan(plan, canonical, store)

    assert _result(outcome, "revenue", "baseline").value == Decimal("0")
    assert _result(outcome, "revenue", "baseline").metric_state is MetricState.VALID
    assert _result(outcome, "revenue", "baseline").currency == "USD"
    assert _record(outcome, "revenue", "baseline").eligible_input_row_count == 0
    assert _record(outcome, "revenue", "baseline").resolved_currency == "USD"
    assert _result(outcome, "orders", "baseline").value == 0
    assert _result(outcome, "orders", "baseline").metric_state is MetricState.VALID
    assert _result(outcome, "aov", "baseline").value is None
    assert _result(outcome, "aov", "baseline").metric_state is MetricState.UNDEFINED
    assert _result(outcome, "aov", "baseline").undefined_reason == "orders_equals_zero"
    assert _result(outcome, "aov", "baseline").currency == "USD"


def test_cancelled_physical_row_with_zero_eligible_rows_remains_valid_empty_population(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(
        tmp_path,
        ("aov",),
        [
            _row(order_id="o1", order_line_id="l1", eligibility_status="cancelled", currency="USD"),
            _row(order_id="o2", order_date="2026-01-03", line_revenue="12.00", currency="USD"),
        ],
        filename="cancelled_zero_eligible.csv",
    )

    outcome = execute_plan(plan, canonical, store)

    assert _result(outcome, "revenue", "baseline").value == Decimal("0")
    assert _result(outcome, "revenue", "baseline").metric_state is MetricState.VALID
    assert _result(outcome, "revenue", "baseline").currency == "USD"
    assert _result(outcome, "orders", "baseline").value == 0
    assert _result(outcome, "orders", "baseline").metric_state is MetricState.VALID
    assert _result(outcome, "aov", "baseline").value is None
    assert _result(outcome, "aov", "baseline").metric_state is MetricState.UNDEFINED
    assert _result(outcome, "aov", "baseline").undefined_reason == "orders_equals_zero"
    assert _result(outcome, "aov", "baseline").currency == "USD"
    assert _record(outcome, "revenue", "baseline").eligible_input_row_count == 0


def test_aov_exact_decimal_division_does_not_use_float(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(
        tmp_path,
        ("aov",),
        [
            _row(order_id="o1", order_line_id="l1", line_revenue="1"),
            _row(order_id="o2", order_line_id="l1", line_revenue="0"),
            _row(order_id="o3", order_line_id="l1", line_revenue="0"),
        ],
    )

    result = _result(execute_plan(plan, canonical, store), "aov", "baseline")

    assert result.value == Decimal("0.33333333333333333333333333333333333333")
    assert isinstance(result.value, Decimal)


def test_aov_decimal_policy_is_independent_of_ambient_decimal_context(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(
        tmp_path,
        ("aov",),
        [
            _row(order_id="o1", order_line_id="l1", line_revenue="1"),
            _row(order_id="o2", order_line_id="l1", line_revenue="0"),
            _row(order_id="o3", order_line_id="l1", line_revenue="0"),
        ],
    )
    expected = Decimal("0.33333333333333333333333333333333333333")
    observed = []

    for precision in (10, 28, 50):
        with localcontext() as context:
            context.prec = precision
            context.rounding = ROUND_HALF_EVEN
            observed.append(_result(execute_plan(plan, canonical, store), "aov", "baseline").value)
    for rounding in (ROUND_DOWN, ROUND_UP):
        with localcontext() as context:
            context.prec = 10
            context.rounding = rounding
            observed.append(_result(execute_plan(plan, canonical, store), "aov", "baseline").value)

    assert observed == [expected] * 5
    assert all(isinstance(value, Decimal) for value in observed)


def test_aov_decimal_policy_does_not_mutate_global_context(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(
        tmp_path,
        ("aov",),
        [
            _row(order_id="o1", order_line_id="l1", line_revenue="1"),
            _row(order_id="o2", order_line_id="l1", line_revenue="0"),
            _row(order_id="o3", order_line_id="l1", line_revenue="0"),
        ],
    )
    context = getcontext()
    original_precision = context.prec
    original_rounding = context.rounding
    context.prec = 10
    context.rounding = ROUND_UP
    try:
        result = _result(execute_plan(plan, canonical, store), "aov", "baseline")
        assert result.value == Decimal("0.33333333333333333333333333333333333333")
        assert result.precision == "p4_aov_decimal_calculation_policy_v1"
        assert result.precision_metadata == {
            "calculation_policy_id": "p4_aov_decimal_calculation_policy_v1",
            "precision": 38,
            "rounding": "ROUND_HALF_EVEN",
        }
        assert context.prec == 10
        assert context.rounding == ROUND_UP
    finally:
        context.prec = original_precision
        context.rounding = original_rounding


def test_executable_revenue_and_orders_execute_while_blocked_nodes_do_not(tmp_path) -> None:
    revenue_plan, canonical, store = _execution_inputs(tmp_path, ("revenue",), [_row()])
    revenue_outcome = execute_plan(revenue_plan, canonical, store)
    assert _result(revenue_outcome, "revenue", "baseline").value == Decimal("10.00")

    blocked_revenue = revenue_plan.model_copy(
        update={
            "ordered_metrics": tuple(
                node.model_copy(
                    update={
                        "planning_state": "blocked",
                        "authorized_requested_metric_refs": (),
                        "failure_details": (_blocked_failure(node),),
                    }
                )
                for node in revenue_plan.ordered_metrics
            ),
            "eligible_requested_metric_refs": (),
            "blocked_metric_refs": ("revenue",),
        }
    )
    assert execute_plan(blocked_revenue, canonical, store).executed_results == ()

    orders_plan, canonical, store = _execution_inputs(tmp_path, ("orders",), [_row(order_id="o2")], filename="orders.csv")
    orders_outcome = execute_plan(orders_plan, canonical, store)
    assert _result(orders_outcome, "orders", "baseline").value == 1

    blocked_orders = orders_plan.model_copy(
        update={
            "ordered_metrics": tuple(
                node.model_copy(
                    update={
                        "planning_state": "blocked",
                        "authorized_requested_metric_refs": (),
                        "failure_details": (_blocked_failure(node),),
                    }
                )
                for node in orders_plan.ordered_metrics
            ),
            "eligible_requested_metric_refs": (),
            "blocked_metric_refs": ("orders",),
        }
    )
    assert execute_plan(blocked_orders, canonical, store).execution_records == ()


def test_blocked_aov_chain_does_not_execute_hidden_dependencies(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(
        tmp_path,
        ("aov",),
        [_row()],
        ineligible={"aov": ()},
        filename="blocked_aov.csv",
    )

    outcome = execute_plan(plan, canonical, store)

    assert outcome.execution_records == ()
    assert outcome.executed_results == ()


def test_aov_malformed_dependency_results_fail_closed(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(tmp_path, ("aov",), [_row()])
    aov_node = next(node for node in plan.ordered_metrics if node.metric_ref == "aov")
    malformed_aov = aov_node.model_copy(update={"dependency_node_ids": aov_node.dependency_node_ids[:1]})
    malformed_plan = plan.model_copy(
        update={
            "ordered_metrics": tuple(
                malformed_aov if node.node_id == aov_node.node_id else node
                for node in plan.ordered_metrics
            ),
        }
    )

    outcome = execute_plan(malformed_plan, canonical, store)
    failed_aov = [record for record in outcome.execution_records if record.metric_refs == ("aov",)]

    assert failed_aov
    assert failed_aov[0].status.value == "failed"
    assert "AOV requires executed Revenue and Orders" in failed_aov[0].failure_details[0].reason
    assert not any(
        result.metric_ref == "aov" and result.period_ref == aov_node.period_refs[0]
        for result in outcome.executed_results
    )


def test_revenue_change_executes_from_baseline_and_comparison_revenue_dependencies(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(
        tmp_path,
        ("revenue_change",),
        [
            _row(order_id="o1", order_date="2026-01-01", line_revenue="100.00"),
            _row(order_id="o2", order_date="2026-01-03", line_revenue="120.00"),
        ],
    )

    outcome = execute_plan(plan, canonical, store)
    result = _result(outcome, "revenue_change", "comparison")
    record = _record(outcome, "revenue_change", "baseline_and_comparison")

    assert result.value == Decimal("20.00")
    assert result.metric_state is MetricState.VALID
    assert result.undefined_reason is None
    assert result.precision == "p7_revenue_change_decimal_calculation_policy_v1"
    assert result.precision_metadata == {
        "calculation_policy_id": "p7_revenue_change_decimal_calculation_policy_v1",
        "precision": 38,
        "rounding": "ROUND_HALF_EVEN",
        "operation": "subtraction",
    }
    assert result.currency == "USD"
    assert record.metric_implementation_ref == REVENUE_CHANGE_EXECUTION_IMPLEMENTATION_REF
    assert record.period_refs == ("baseline", "comparison")
    assert record.period_role == "baseline_and_comparison"
    assert record.population_refs != ()
    assert record.operation["formula"] == "comparison_revenue - baseline_revenue"
    assert "baseline_revenue_result_ref" in record.operation
    assert "comparison_revenue_result_ref" in record.operation


@pytest.mark.parametrize(
    ("tamper", "expected_reason"),
    (
        ("wrong_request", "request authority mismatches plan"),
        ("wrong_plan", "plan authority mismatches plan"),
        ("wrong_node", "does not match governed dependency node"),
        ("fabricated_result", "result_ref mismatches result"),
        ("missing_artifact", "artifact integrity check failed"),
        ("tampered_artifact", "artifact integrity check failed"),
    ),
)
def test_revenue_change_execution_dependency_authority_fails_closed(tmp_path, tamper: str, expected_reason: str) -> None:
    plan, canonical, store = _execution_inputs(
        tmp_path,
        ("revenue_change",),
        [
            _row(order_id="o1", order_date="2026-01-01", line_revenue="100.00"),
            _row(order_id="o2", order_date="2026-01-03", line_revenue="120.00"),
        ],
    )
    metadata_store = MetadataStore(tmp_path / f"{tamper}_execution_registry.sqlite")
    outcome = execute_plan(plan, canonical, store, metadata_store)
    change_node = next(node for node in plan.ordered_metrics if node.metric_ref == "revenue_change")
    population_by_id = {population.population_id: population for population in plan.population_definitions}
    result_by_node_id = {
        node.node_id: _result(outcome, "revenue", node.period_refs[0])
        for node in plan.ordered_metrics
        if node.metric_ref == "revenue"
    }
    baseline_result = _result(outcome, "revenue", "baseline")
    baseline_record = metadata_store.get_execution_record(baseline_result.execution_id)
    assert baseline_record is not None

    if tamper == "wrong_request":
        _replace_execution_record(metadata_store, baseline_record.model_copy(update={"request_id": "req_wrong"}))
    elif tamper == "wrong_plan":
        _replace_execution_record(metadata_store, baseline_record.model_copy(update={"plan_id": "plan_wrong"}))
    elif tamper == "wrong_node":
        _replace_execution_record(metadata_store, baseline_record.model_copy(update={"plan_node_id": "node_wrong"}))
    elif tamper == "fabricated_result":
        result_by_node_id[baseline_record.plan_node_id] = baseline_result.model_copy(update={"result_id": "exres_fabricated"})
    elif tamper == "missing_artifact":
        store.safe_path(baseline_record.output_artifacts[0].path).unlink()
    elif tamper == "tampered_artifact":
        store.safe_path(baseline_record.output_artifacts[0].path).write_text("tampered", encoding="utf-8")

    with pytest.raises(MetricExecutionError, match=expected_reason):
        _revenue_change_dependency_results(
            plan,
            change_node,
            result_by_node_id,
            population_by_id,
            metadata_store,
            store,
        )


@pytest.mark.parametrize(
    ("replacement_ref", "expected_reason"),
    (
        (EXECUTION_NOT_IMPLEMENTED_REF, "does not match approved Metric Registry binding"),
        ("review:mismatched_implementation_ref", "does not match approved Metric Registry binding"),
        (ORDERS_EXECUTION_IMPLEMENTATION_REF, "does not match approved Metric Registry binding"),
    ),
)
def test_executable_metric_requires_registry_approved_implementation_ref(
    tmp_path,
    replacement_ref: str,
    expected_reason: str,
) -> None:
    plan, canonical, store = _execution_inputs(tmp_path, ("revenue",), [_row()])
    forged_node = plan.ordered_metrics[0].model_copy(update={"execution_implementation_ref": replacement_ref})
    forged_plan = plan.model_copy(update={"ordered_metrics": (forged_node, *plan.ordered_metrics[1:])})

    outcome = execute_plan(forged_plan, canonical, store)
    failed = _record(outcome, "revenue", "baseline")

    assert failed.status.value == "failed"
    assert expected_reason in failed.failure_details[0].reason
    assert failed.started_at <= failed.ended_at
    assert not any(result.metric_ref == "revenue" and result.period_ref == "baseline" for result in outcome.executed_results)
    assert _result(outcome, "revenue", "comparison").value == Decimal("0")


@pytest.mark.parametrize(
    ("baseline", "comparison", "expected"),
    [
        ("100.00", "120.00", Decimal("20.00")),
        ("120.00", "100.00", Decimal("-20.00")),
        ("100.00", "100.00", Decimal("0.00")),
        ("0.00", "100.00", Decimal("100.00")),
        ("100.00", "0.00", Decimal("-100.00")),
        ("0.00", "0.00", Decimal("0.00")),
    ],
)
def test_revenue_change_positive_negative_zero_and_empty_period_cases_are_valid(
    tmp_path,
    baseline: str,
    comparison: str,
    expected: Decimal,
) -> None:
    rows = []
    if Decimal(baseline) != 0:
        rows.append(_row(order_id="o1", order_date="2026-01-01", line_revenue=baseline))
    if Decimal(comparison) != 0:
        rows.append(_row(order_id="o2", order_date="2026-01-03", line_revenue=comparison))
    if not rows:
        rows.append(_row(order_id="o3", order_date="2026-01-03", line_revenue="0.00", eligibility_status="cancelled"))
    plan, canonical, store = _execution_inputs(tmp_path, ("revenue_change",), rows, filename=f"change_{baseline}_{comparison}.csv")

    outcome = execute_plan(plan, canonical, store)
    result = _result(outcome, "revenue_change", "comparison")

    assert result.value == expected
    assert result.metric_state is MetricState.VALID
    assert result.undefined_reason is None


def test_supported_dependencies_and_revenue_change_root_all_execute(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(
        tmp_path,
        ("revenue_change",),
        [
            _row(order_id="o1", order_date="2026-01-01", line_revenue="10.00"),
            _row(order_id="o2", order_date="2026-01-03", line_revenue="11.00"),
        ],
        filename="partial_revenue_change.csv",
    )

    outcome = execute_plan(plan, canonical, store)

    assert _result(outcome, "revenue", "baseline").value == Decimal("10.00")
    assert _result(outcome, "revenue", "comparison").value == Decimal("11.00")
    assert _result(outcome, "revenue_change", "comparison").value == Decimal("1.00")
    assert _record(outcome, "revenue_change", "baseline_and_comparison").status.value == "completed"


def test_malformed_executable_authorization_fails_before_execution(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(tmp_path, ("revenue",), [_row()])
    bad_node = plan.ordered_metrics[0].model_copy(update={"authorized_requested_metric_refs": ()})
    bad_plan = plan.model_copy(update={"ordered_metrics": (bad_node, *plan.ordered_metrics[1:])})

    with pytest.raises(Exception):
        execute_plan(bad_plan, canonical, store)


def test_provenance_records_plan_node_metric_dataset_population_executor_operation_and_result(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(
        tmp_path,
        ("revenue",),
        [_row()],
        scope=ScopeDefinition(
            scope_id="usd_only",
            filters=(ScopeFilter(field="currency", operator="equals", value="USD"),),
        ),
    )

    outcome = execute_plan(plan, canonical, store)
    record = _record(outcome, "revenue", "baseline")

    assert record.plan_id == plan.plan_id
    assert record.plan_fingerprint == plan.plan_fingerprint
    assert record.plan_node_id
    assert record.metric_refs == ("revenue",)
    assert record.metric_definition_version == METRIC_DEFINITION_VERSION
    assert record.metric_implementation_ref == REVENUE_EXECUTION_IMPLEMENTATION_REF
    assert record.canonical_dataset_ref_ids == (canonical.canonical_dataset_id,)
    assert record.canonical_dataset_fingerprints == (canonical.content_fingerprint,)
    assert record.population_refs
    assert record.population_fingerprints
    assert record.periods[0]["period_id"] == "baseline"
    assert record.scope_filters == ({"field": "currency", "operator": "equals", "value": "USD"},)
    assert record.grouping == "none"
    assert record.resolved_currency == "USD"
    assert record.eligible_input_row_count == 1
    assert record.executor_id == "commerce_lens_duckdb_reference_executor"
    assert record.executor_version == "p4_001_v1"
    assert record.duckdb_version
    assert record.operation["method"] == "duckdb_sql"
    assert "SUM(line_revenue)" in record.operation["sql"]
    assert record.result_ref == _result(outcome, "revenue", "baseline").result_id
    assert record.output_artifacts
    assert not hasattr(record, "validation_passed")
    assert not hasattr(record, "admissible_evidence_ref")


def test_mixed_runtime_currency_fails_closed_for_monetary_metrics(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(
        tmp_path,
        ("revenue",),
        [
            _row(order_id="o1", line_revenue="10.00", currency="USD"),
            _row(order_id="o2", order_line_id="l1", line_revenue="20.00", currency="USD"),
        ],
        filename="mixed_runtime_currency.csv",
    )
    tampered = _replace_canonical_parquet_rows(
        canonical,
        store,
        "SELECT * REPLACE (CASE WHEN order_id = 'o2' THEN 'EUR' ELSE currency END AS currency) FROM read_parquet(?)",
    )

    outcome = execute_plan(plan, tampered, store)
    failed = _record(outcome, "revenue", "baseline")

    assert failed.status.value == "failed"
    assert "single-governed-currency" in failed.failure_details[0].reason
    assert not any(result.metric_ref == "revenue" and result.period_ref == "baseline" for result in outcome.executed_results)


def test_explicit_currency_with_contradictory_participating_currency_fails_closed(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(
        tmp_path,
        ("revenue",),
        [_row(order_id="o1", line_revenue="10.00", currency="EUR")],
        filename="explicit_currency_contradiction.csv",
    )
    usd_plan = _plan_with_population_currency_basis(plan, "baseline", "currency:USD")

    outcome = execute_plan(usd_plan, canonical, store)
    failed = _record(outcome, "revenue", "baseline")

    assert failed.status.value == "failed"
    assert "does not match governed currency scope filter" in failed.failure_details[0].reason
    assert not any(result.metric_ref == "revenue" and result.period_ref == "baseline" for result in outcome.executed_results)


def test_phase2_single_currency_with_mixed_dataset_authority_fails_closed_for_empty_period(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(
        tmp_path,
        ("revenue",),
        [
            _row(order_id="o1", order_date="2026-01-03", line_revenue="10.00", currency="USD"),
            _row(order_id="o2", order_date="2026-01-04", line_revenue="20.00", currency="USD"),
        ],
        filename="phase2_mixed_dataset_authority.csv",
    )
    tampered = _replace_canonical_parquet_rows(
        canonical,
        store,
        "SELECT * REPLACE (CASE WHEN order_id = 'o2' THEN 'EUR' ELSE currency END AS currency) FROM read_parquet(?)",
    )

    outcome = execute_plan(plan, tampered, store)
    failed = _record(outcome, "revenue", "baseline")

    assert failed.status.value == "failed"
    assert "single-governed-currency" in failed.failure_details[0].reason
    assert not any(result.metric_ref == "revenue" and result.period_ref == "baseline" for result in outcome.executed_results)


def test_phase2_single_currency_with_no_currency_evidence_fails_closed(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(
        tmp_path,
        ("revenue",),
        [_row(order_id="o1", line_revenue="10.00", currency="USD")],
        filename="empty_currency_unresolvable.csv",
    )
    empty_canonical = _replace_canonical_parquet_rows(
        canonical,
        store,
        "SELECT * FROM read_parquet(?) WHERE false",
    ).model_copy(update={"row_count": 0})

    outcome = execute_plan(plan, empty_canonical, store)
    failed = _record(outcome, "revenue", "baseline")

    assert failed.status.value == "failed"
    assert "no canonical rows from which to establish phase2 single-governed currency authority" in (
        failed.failure_details[0].reason
    )
    assert not any(result.metric_ref == "revenue" and result.period_ref == "baseline" for result in outcome.executed_results)


def test_stale_population_fingerprint_fails_closed_before_execution(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(tmp_path, ("revenue",), [_row()])
    baseline_population = next(pop for pop in plan.population_definitions if pop.period.period_id == "baseline")
    tampered_period = baseline_population.period.model_copy(update={"end_date": date(2026, 1, 3)})
    stale_population = baseline_population.model_copy(update={"period": tampered_period})
    stale_plan = plan.model_copy(
        update={
            "population_definitions": tuple(
                stale_population if pop.population_id == baseline_population.population_id else pop
                for pop in plan.population_definitions
            )
        }
    )

    outcome = execute_plan(stale_plan, canonical, store)
    failed = _record(outcome, "revenue", "baseline")

    assert failed.status.value == "failed"
    assert "population fingerprint" in failed.failure_details[0].reason
    assert not any(result.metric_ref == "revenue" and result.period_ref == "baseline" for result in outcome.executed_results)


def test_completed_execution_record_and_result_artifact_persist(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(tmp_path, ("revenue",), [_row()])
    metadata_store = MetadataStore(tmp_path / "execution_registry.sqlite")

    outcome = execute_plan(plan, canonical, store, metadata_store)
    record = _record(outcome, "revenue", "baseline")
    result = _result(outcome, "revenue", "baseline")
    stored_record = metadata_store.get_execution_record(record.execution_id)

    assert stored_record == record
    assert stored_record.output_artifacts
    artifact = metadata_store.get_artifact_reference(stored_record.output_artifacts[0].artifact_id)
    assert artifact == stored_record.output_artifacts[0]
    artifact_path = store.safe_path(artifact.path)
    restored = ExecutedResult.model_validate_json(artifact_path.read_text(encoding="utf-8"))
    assert restored == result
    assert restored.value == Decimal("10.00")
    assert isinstance(restored.value, Decimal)
    assert restored.execution_id == record.execution_id


def test_undefined_aov_result_artifact_persists(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(
        tmp_path,
        ("aov",),
        [_row(order_id="o1", order_line_id="l1", eligibility_status="cancelled")],
        filename="undefined_aov_persistence.csv",
    )
    metadata_store = MetadataStore(tmp_path / "undefined_registry.sqlite")

    outcome = execute_plan(plan, canonical, store, metadata_store)
    result = _result(outcome, "aov", "baseline")
    record = _record(outcome, "aov", "baseline")
    artifact_path = store.safe_path(record.output_artifacts[0].path)
    restored = ExecutedResult.model_validate_json(artifact_path.read_text(encoding="utf-8"))

    assert restored.metric_state is MetricState.UNDEFINED
    assert restored.undefined_reason == "orders_equals_zero"
    assert restored.value is None
    assert restored.precision_metadata["calculation_policy_id"] == "p4_aov_decimal_calculation_policy_v1"
    assert restored == result


def test_revenue_change_execution_record_and_artifact_persist(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(tmp_path, ("revenue_change",), [_row()])
    metadata_store = MetadataStore(tmp_path / "failed_registry.sqlite")

    outcome = execute_plan(plan, canonical, store, metadata_store)
    record = _record(outcome, "revenue_change", "baseline_and_comparison")
    result = _result(outcome, "revenue_change", "comparison")
    stored_record = metadata_store.get_execution_record(record.execution_id)

    assert stored_record == record
    assert stored_record.status.value == "completed"
    assert stored_record.result_ref == result.result_id
    assert stored_record.output_artifacts
    restored = ExecutedResult.model_validate_json(
        store.safe_path(stored_record.output_artifacts[0].path).read_text(encoding="utf-8")
    )
    assert restored == result


def test_repeated_equivalent_execution_attempts_persist_separately(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(tmp_path, ("revenue",), [_row()])
    metadata_store = MetadataStore(tmp_path / "repeat_registry.sqlite")

    first = execute_plan(plan, canonical, store, metadata_store)
    second = execute_plan(plan, canonical, store, metadata_store)
    records = [
        record
        for record in metadata_store.list_execution_records()
        if record.metric_refs == ("revenue",) and record.period_refs == ("baseline",)
    ]

    assert len(records) == 2
    assert {record.execution_id for record in records} == {
        _record(first, "revenue", "baseline").execution_id,
        _record(second, "revenue", "baseline").execution_id,
    }
    assert records[0].execution_id != records[1].execution_id


def test_missing_or_mismatched_canonical_artifact_fails_closed(tmp_path) -> None:
    plan, canonical, store = _execution_inputs(tmp_path, ("revenue",), [_row()])
    missing = canonical.model_copy(
        update={"artifact": canonical.artifact.model_copy(update={"path": "canonical/missing.parquet"})}
    )

    with pytest.raises(MetricExecutionError):
        execute_plan(plan, missing, store)


def test_csv_xlsx_sqlite_converge_through_canonical_artifact_execution(tmp_path) -> None:
    rows = [
        _row(order_id="o1", line_revenue="10.25"),
        _row(order_id="o2", order_date="2026-01-03", line_revenue="20.50"),
    ]
    outcomes = []
    for source_type in ("csv", "xlsx", "sqlite"):
        plan, canonical, store = _execution_inputs(
            tmp_path,
            ("revenue", "orders", "aov"),
            rows,
            filename=f"converge.{source_type}",
            source_type=source_type,
        )
        outcomes.append(execute_plan(plan, canonical, store))

    observed = [
        (
            _result(outcome, "revenue", "baseline").value,
            _result(outcome, "orders", "baseline").value,
            _result(outcome, "aov", "comparison").value,
        )
        for outcome in outcomes
    ]
    assert observed == [(Decimal("10.25"), 1, Decimal("20.50"))] * 3


def _execution_inputs(
    tmp_path: Path,
    metrics: tuple[str, ...],
    rows: list[dict[str, str]],
    *,
    ineligible: dict[str, tuple[FailureDetail, ...]] | None = None,
    filename: str = "orders.csv",
    source_type: str = "csv",
    scope: ScopeDefinition = ScopeDefinition(scope_id="all_eligible"),
):
    dataset, store = _registered_source(tmp_path, rows, filename=filename, source_type=source_type)
    canonicalization = canonicalize_dataset(dataset, _canonicalization_request(dataset), store)
    assert canonicalization.canonical_dataset is not None
    request = _request(metrics, dataset.dataset_id, scope=scope)
    plan = build_execution_plan(
        request,
        _sufficiency(
            request,
            canonicalization.canonical_dataset.canonical_dataset_id,
            ineligible=ineligible,
        ),
    )
    return plan, canonicalization.canonical_dataset, store


def _request(metrics: tuple[str, ...], dataset_ref_id: str, *, scope: ScopeDefinition) -> AnalysisRequest:
    return AnalysisRequest(
        canonical_business_question_id="canonical_revenue_change",
        metrics=tuple(
            MetricReference(metric_id=metric_id, definition_version=METRIC_DEFINITION_VERSION)
            for metric_id in metrics
        ),
        baseline_period=PeriodDefinition(
            period_id="baseline",
            label="Baseline",
            start_date=date(2026, 1, 1),
            end_date=date(2026, 1, 2),
            date_convention_ref="order_date_utc",
        ),
        comparison_period=PeriodDefinition(
            period_id="comparison",
            label="Comparison",
            start_date=date(2026, 1, 3),
            end_date=date(2026, 1, 4),
            date_convention_ref="order_date_utc",
        ),
        scope=scope,
        grouping=GroupingDimension.NONE,
        dataset_ref_id=dataset_ref_id,
        canonical_schema_version="canonical_mvp_v1",
        metric_registry_version=METRIC_REGISTRY_VERSION,
    )


def _sufficiency(
    request: AnalysisRequest,
    canonical_dataset_ref_id: str,
    *,
    ineligible: dict[str, tuple[FailureDetail, ...]] | None = None,
) -> DataSufficiencyResult:
    ineligible = ineligible or {}
    return DataSufficiencyResult(
        sufficiency_id=f"suff_{request.request_id}",
        request_id=request.request_id,
        dataset_ref_id=request.dataset_ref_id,
        canonical_dataset_ref_id=canonical_dataset_ref_id,
        metric_eligibility=tuple(
            MetricEligibility(
                metric_ref=metric.metric_id,
                eligible=metric.metric_id not in ineligible,
                metric_state=MetricState.INADMISSIBLE if metric.metric_id in ineligible else None,
                failure_details=ineligible.get(metric.metric_id, ()),
            )
            for metric in request.metrics
        ),
        state=SufficiencyState.PARTIAL if ineligible else SufficiencyState.SUFFICIENT,
    )


def _canonicalization_request(dataset) -> CanonicalizationRequest:
    headers = _headers()
    return CanonicalizationRequest(
        source_dataset_id=dataset.dataset_id,
        mapping=identity_mapping(headers, require_eligibility=True),
        eligibility_mode=EligibilityMode.EXPLICIT_STATUS_MAPPING,
        eligibility_value_mapping=(
            EligibilityValueMapping(source_value="paid", normalized_status=EligibilityState.ELIGIBLE),
            EligibilityValueMapping(source_value="cancelled", normalized_status=EligibilityState.EXCLUDED),
        ),
    )


def _registered_source(
    tmp_path: Path,
    rows: list[dict[str, str]],
    *,
    filename: str,
    source_type: str,
):
    path = tmp_path / filename
    if source_type == "csv":
        _write_csv(path, rows)
        selected_sheet = None
        selected_table = None
        source = SourceType.CSV
    elif source_type == "xlsx":
        _write_xlsx(path, rows)
        selected_sheet = "Orders"
        selected_table = None
        source = SourceType.EXCEL_XLSX
    elif source_type == "sqlite":
        _write_sqlite(path, rows)
        selected_sheet = None
        selected_table = "orders"
        source = SourceType.SQLITE
    else:
        raise AssertionError(f"unknown source_type {source_type}")
    store = ArtifactStore(tmp_path / f"runtime_{filename}")
    registry = DatasetRegistry(store)
    return registry.register_source(path, source, selected_sheet=selected_sheet, selected_table=selected_table), store


def _write_csv(path: Path, rows: list[dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as file_obj:
        writer = csv.DictWriter(file_obj, fieldnames=_headers())
        writer.writeheader()
        writer.writerows([{key: row.get(key, "") for key in _headers()} for row in rows])


def _write_xlsx(path: Path, rows: list[dict[str, str]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Orders"
    sheet.append(_headers())
    for row in rows:
        sheet.append([row.get(key, "") for key in _headers()])
    workbook.save(path)


def _write_sqlite(path: Path, rows: list[dict[str, str]]) -> None:
    conn = sqlite3.connect(path)
    with conn:
        columns = ", ".join(f"{header} TEXT" for header in _headers())
        conn.execute(f"CREATE TABLE orders ({columns})")
        placeholders = ", ".join("?" for _ in _headers())
        conn.executemany(
            f"INSERT INTO orders VALUES ({placeholders})",
            [[row.get(key, "") for key in _headers()] for row in rows],
        )
    conn.close()


def _replace_canonical_parquet_rows(canonical, store: ArtifactStore, select_sql: str):
    path = store.safe_path(canonical.artifact.path)
    replacement = path.with_name(f"{path.stem}_replacement.parquet")
    input_sql = str(path).replace("'", "''")
    output_sql = str(replacement).replace("'", "''")
    conn = duckdb.connect(database=":memory:")
    try:
        conn.execute(f"COPY ({select_sql}) TO '{output_sql}' (FORMAT PARQUET)", (input_sql,))
    finally:
        conn.close()
    path.unlink()
    replacement.rename(path)
    fingerprint = sha256_file(path)
    return canonical.model_copy(
        update={
            "content_fingerprint": fingerprint,
            "artifact": canonical.artifact.model_copy(
                update={"fingerprint": fingerprint, "size_bytes": path.stat().st_size}
            ),
        }
    )


def _replace_execution_record(metadata_store: MetadataStore, record) -> None:
    result_artifact_id = record.output_artifacts[0].artifact_id if record.output_artifacts else None
    with sqlite3.connect(metadata_store.db_path) as conn:
        conn.execute(
            """
            UPDATE execution_records
            SET request_id = ?,
                plan_id = ?,
                plan_node_id = ?,
                metric_ref = ?,
                status = ?,
                result_ref = ?,
                result_artifact_id = ?,
                started_at = ?,
                ended_at = ?,
                record_json = ?
            WHERE execution_id = ?
            """,
            (
                record.request_id,
                record.plan_id,
                record.plan_node_id,
                record.metric_refs[0] if record.metric_refs else None,
                record.status.value,
                record.result_ref,
                result_artifact_id,
                record.started_at.isoformat(),
                record.ended_at.isoformat() if record.ended_at is not None else None,
                record.model_dump_json(),
                record.execution_id,
            ),
        )


def _plan_with_population_currency_basis(plan, period_ref: str, currency_basis_ref: str):
    old_population = next(
        population
        for population in plan.population_definitions
        if population.period.period_id == period_ref and population.grouping is GroupingDimension.NONE
    )
    draft_population = old_population.model_copy(update={"currency_basis_ref": currency_basis_ref})
    fingerprint = population_fingerprint(draft_population)
    new_population = draft_population.model_copy(
        update={
            "population_fingerprint": fingerprint,
            "population_id": population_id_for_fingerprint(fingerprint),
        }
    )
    return plan.model_copy(
        update={
            "ordered_metrics": tuple(
                node.model_copy(
                    update={
                        "population_refs": tuple(
                            new_population.population_id if ref == old_population.population_id else ref
                            for ref in node.population_refs
                        )
                    }
                )
                for node in plan.ordered_metrics
            ),
            "population_definitions": tuple(
                new_population if population.population_id == old_population.population_id else population
                for population in plan.population_definitions
            ),
            "population_refs": tuple(
                new_population.population_id if ref == old_population.population_id else ref
                for ref in plan.population_refs
            ),
        }
    )


def _headers() -> tuple[str, ...]:
    return (
        "order_id",
        "order_line_id",
        "order_date",
        "product_id",
        "product_name",
        "category_id",
        "category_name",
        "quantity",
        "line_revenue",
        "currency",
        "unit_price",
        "eligibility_status",
    )


def _row(**overrides: str) -> dict[str, str]:
    row = {
        "order_id": "o1",
        "order_line_id": "l1",
        "order_date": "2026-01-01",
        "product_id": "p1",
        "product_name": "Tea",
        "category_id": "c1",
        "category_name": "Drinks",
        "quantity": "1",
        "line_revenue": "10.00",
        "currency": "USD",
        "unit_price": "",
        "eligibility_status": "paid",
    }
    row.update(overrides)
    return row


def _result(outcome, metric_ref: str, period_ref: str):
    return next(
        result
        for result in outcome.executed_results
        if result.metric_ref == metric_ref and result.period_ref == period_ref
    )


def _record(outcome, metric_ref: str, period_ref: str):
    return next(
        record
        for record in outcome.execution_records
        if record.metric_refs == (metric_ref,)
        and (
            record.period_refs == (period_ref,)
            or (
                metric_ref == "revenue_change"
                and period_ref in {"comparison", "baseline_and_comparison"}
                and record.period_refs == ("baseline", "comparison")
            )
        )
    )


def _blocked_failure(node: PlanMetricNode) -> FailureDetail:
    return FailureDetail(
        stage=FailureStage.PLANNING,
        reason="test blocked by Data Sufficiency",
        target_ref=node.node_id,
        governing_ref="tasks:P3-001",
        dependency_scope=node.metric_ref,
        independent_chains_may_continue=True,
    )
